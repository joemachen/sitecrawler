from pathlib import Path

from openpyxl import load_workbook

from sitecrawler.export import write_xlsx
from sitecrawler.store.dump import TableDump


def test_one_sheet_per_table_with_headers(tmp_path: Path) -> None:
    tables = [
        TableDump(
            "pages",
            ("url", "status_code", "ok"),
            [("https://x/", 200, True), ("https://x/a", None, False)],
        ),
        TableDump("links", ("src", "dst"), []),
    ]
    out = tmp_path / "crawl.xlsx"
    assert write_xlsx(tables, out) == 2

    workbook = load_workbook(out, read_only=True)
    try:
        assert workbook.sheetnames == ["pages", "links"]
        rows = list(workbook["pages"].iter_rows(values_only=True))
        assert rows[0] == ("url", "status_code", "ok")
        assert rows[1] == ("https://x/", 200, True)
        assert rows[2] == ("https://x/a", None, False)
        assert list(workbook["links"].iter_rows(values_only=True)) == [("src", "dst")]
    finally:
        workbook.close()
